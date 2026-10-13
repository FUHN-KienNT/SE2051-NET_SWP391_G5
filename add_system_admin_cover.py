import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

FILE_PATH = 'doc/Integration-Test.xlsx'

wb = openpyxl.load_workbook(FILE_PATH)
ws = wb['Cover']

FONT_TAHOMA_9 = Font(name="Tahoma", size=9, bold=False, color="000000")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_SIDE = Side(border_style="thin", color="000000")
BORDER_THIN_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

# Add System Administration change log at row 13
row = 13
ws.row_dimensions[row].height = 42.0

data = [
    ("2026-10-04", ALIGN_CENTER),
    ("v1.0", ALIGN_CENTER),
    ("System Administration", ALIGN_CENTER),
    ("A", ALIGN_CENTER),
    ("Add integration test suite (21 TCs) for System Administration across Course Management screens: Course Search, Multi-criteria Filter, Course Detail editing, Role Assignment (Manager/Expert), and Course Lifecycle Status updates. Executed Round 1 & Round 2.", ALIGN_LEFT),
    ("Software Requirement Specification (SRS II.2), Software Design Specification (SDS)", ALIGN_LEFT)
]

for col_idx, (val, align) in enumerate(data, 1):
    c = ws.cell(row, col_idx, val)
    c.font = FONT_TAHOMA_9
    c.alignment = align
    c.border = BORDER_THIN_ALL

wb.save(FILE_PATH)
print("Successfully added System Administration to Cover sheet!")
