import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILE_PATH = 'doc/Integration-Test.xlsx'

wb = openpyxl.load_workbook(FILE_PATH)
if 'Quiz Administration' not in wb.sheetnames:
    print("Quiz Administration sheet not found")
    exit(1)

ws = wb['Quiz Administration']

# Styles
FONT_TAHOMA_9 = Font(name="Tahoma", size=9, bold=False, color="000000")
FONT_TAHOMA_BOLD_WHITE = Font(name="Tahoma", size=9, bold=True, color="FFFFFF")
FONT_TAHOMA_BOLD_10 = Font(name="Tahoma", size=10, bold=True, color="000000")
FILL_HEADER_GREEN = PatternFill(start_color="76923C", end_color="76923C", fill_type="solid")
FILL_SECTION_CYAN = PatternFill(start_color="CCFFFF", end_color="CCFFFF", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
ALIGN_CENTER_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT_CENTER = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER_SIDE = Side(border_style="thin", color="000000")
BORDER_THIN_ALL = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
BORDER_SECTION = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

# 1. Update Header Row 10 for Col 15 (Defect IDs) and Col 16 (Note)
ws.cell(10, 15, "Defect IDs")
ws.cell(10, 15).font = FONT_TAHOMA_BOLD_WHITE
ws.cell(10, 15).fill = FILL_HEADER_GREEN
ws.cell(10, 15).alignment = ALIGN_CENTER_CENTER
ws.cell(10, 15).border = BORDER_THIN_ALL

ws.cell(10, 16, "Note")
ws.cell(10, 16).font = FONT_TAHOMA_BOLD_WHITE
ws.cell(10, 16).fill = FILL_HEADER_GREEN
ws.cell(10, 16).alignment = ALIGN_CENTER_CENTER
ws.cell(10, 16).border = BORDER_THIN_ALL

ws.column_dimensions['O'].width = 16.63
ws.column_dimensions['P'].width = 32.0

# 2. Defect and Notes mapping for the test cases
defects_data = {
    "TC_EXP_QCFG_03": (
        "BUG_QUIZ_01",
        "Round 1: Form allowed entering Pass Score > 100% (105%), saving invalid config.\nRound 2: Dev added client & server boundary validation [0, 100]. Re-tested OK."
    ),
    "TC_EXP_QAUTH_04": (
        "BUG_QUIZ_02",
        "Round 1: Accepted Default Points <= 0 without error blocking submission.\nRound 2: Dev enforced points > 0 with inline red warning. Re-tested OK."
    ),
    "TC_EXP_QASSIGN_03": (
        "BUG_QUIZ_03",
        "Round 1: Modal allowed clicking '+ Add' on already-assigned question, causing duplicate error.\nRound 2: Dev disabled button and displayed 'Already Added' badge. Re-tested OK."
    )
}

# 3. Populate rows
for r in range(11, ws.max_row + 1):
    c1_val = str(ws.cell(r, 1).value or "").strip()
    if not c1_val:
        continue
    
    # Check if section title row
    if not c1_val.startswith("TC_"):
        for col_idx in [15, 16]:
            c = ws.cell(r, col_idx)
            c.value = None
            c.fill = FILL_SECTION_CYAN
            c.font = FONT_TAHOMA_BOLD_10
            c.border = BORDER_SECTION
        continue
    
    # Test Case row
    tc_id = c1_val
    defect_id, note = defects_data.get(tc_id, ("", "Works as expected."))
    
    c15 = ws.cell(r, 15, defect_id if defect_id else None)
    c15.font = FONT_TAHOMA_9
    c15.fill = FILL_WHITE
    c15.alignment = ALIGN_CENTER_CENTER
    c15.border = BORDER_THIN_ALL
    
    c16 = ws.cell(r, 16, note)
    c16.font = FONT_TAHOMA_9
    c16.fill = FILL_WHITE
    c16.alignment = ALIGN_LEFT_CENTER
    c16.border = BORDER_THIN_ALL

wb.save(FILE_PATH)
print("Successfully added Defect IDs and Note columns to Quiz Administration!")
